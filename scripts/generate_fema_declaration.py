#!/usr/bin/env python3
"""
FEMA 23(R) / HDFC EDF Declaration Generator
Fills the official HDFC Bank "Request letter for Export of Services & EDF Filing Cum Disposal Instructions for Credit"
Excel template for any foreign currency invoice.
"""

import sys
import os
import json
import openpyxl

def number_to_words(n, currency=""):
    units = ['', 'One', 'Two', 'Three', 'Four', 'Five', 'Six', 'Seven', 'Eight', 'Nine', 'Ten',
             'Eleven', 'Twelve', 'Thirteen', 'Fourteen', 'Fifteen', 'Sixteen', 'Seventeen', 'Eighteen', 'Nineteen']
    tens = ['', '', 'Twenty', 'Thirty', 'Forty', 'Fifty', 'Sixty', 'Seventy', 'Eighty', 'Ninety']
    
    def helper(num):
        if num == 0:
            return ''
        elif num < 20:
            return units[num] + ' '
        elif num < 100:
            return tens[num // 10] + (' ' + units[num % 10] if num % 10 != 0 else '') + ' '
        elif num < 1000:
            return units[num // 100] + ' Hundred ' + helper(num % 100)
        elif num < 1000000:
            return helper(num // 1000) + 'Thousand ' + helper(num % 1000)
        elif num < 1000000000:
            return helper(num // 1000000) + 'Million ' + helper(num % 1000000)
        else:
            return helper(num // 1000000000) + 'Billion ' + helper(num % 1000000000)
    
    val = float(n)
    whole = int(val)
    frac = int(round((val - whole) * 100))
    res = helper(whole).strip()
    if not res:
        res = 'Zero'
        
    curr_str = f" {currency.upper()}" if currency else ""
    if frac > 0:
        res += f"{curr_str} and {helper(frac).strip()} Cents Only"
    else:
        res += f"{curr_str} Only"
        
    return res

def main():
    if len(sys.argv) < 3:
        print("Usage: generate_fema_declaration.py <input_json> <output_xlsx> [template_xlsx]", file=sys.stderr)
        sys.exit(1)

    input_json_path = sys.argv[1]
    output_xlsx_path = sys.argv[2]
    
    if len(sys.argv) >= 4 and os.path.exists(sys.argv[3]):
        template_path = sys.argv[3]
    else:
        base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
        template_path = os.path.join(base_dir, 'storage', 'app', 'templates', 'fema_23r_services_declaration_template.xlsx')

    if not os.path.exists(template_path):
        print(f"Template not found at: {template_path}", file=sys.stderr)
        sys.exit(2)

    with open(input_json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    wb = openpyxl.load_workbook(template_path)
    
    currency = data.get('currency_code', 'USD')
    grand_total = float(data.get('grand_total', 0.0))
    amount_in_words = number_to_words(grand_total, currency)
    
    exporter = data.get('exporter', {})
    client = data.get('client', {})
    items = data.get('items', [])
    
    if not items:
        items = [{
            'name': data.get('services_description', 'Software & IT Consultancy Services'),
            'sac_code': data.get('sac_code', '998314'),
            'amount': grand_total
        }]
    
    # =========================================================================
    # 1. SHEET: REQUEST LETTER
    # =========================================================================
    if 'REQUEST LETTER' in wb.sheetnames:
        ws_req = wb['REQUEST LETTER']
        
        # Letter Date
        letter_date = data.get('request_letter_date') or data.get('payment_date') or data.get('invoice_date')
        if letter_date:
            ws_req['AG4'] = letter_date
            
        # Account to be Credited
        account_no = exporter.get('account_number', '50200120756905')
        ws_req['H5'] = account_no
        
        # Purpose Code
        purpose_code = data.get('purpose_code', 'P0802 - Software consultancy implementation other than those covered in SOFTEX form')
        ws_req['F6'] = purpose_code
        
        # Exchange Rate / Forward contract details
        rate_payment = data.get('exchange_rate_payment')
        rate_invoice = data.get('exchange_rate_invoice')
        if rate_payment:
            ws_req['L7'] = f"1 {currency} = {float(rate_payment):.4f} INR (Realized Bank Liquidation Rate)"
        elif rate_invoice:
            ws_req['L7'] = f"1 {currency} = {float(rate_invoice):.4f} INR (Rate on Invoice Date)"
        else:
            ws_req['L7'] = "As per Bank Realisation"
            
        # IEC, GSTIN, PAN
        iec = exporter.get('iec', 'N/A')
        gstin = exporter.get('gstin', '')
        pan = exporter.get('pan', '')
        
        ws_req['P8'] = f"IEC: {iec}" if iec else "IEC: N/A"
        ws_req['AB8'] = f"GSTIN: {gstin}" if gstin else "GSTIN: N/A"
        ws_req['A9'] = f"Customer PAN: {pan}" if pan else "Customer PAN: N/A"
        
        # Bank & Exporter Details
        ws_req['P10'] = exporter.get('ad_name_address') or f"{exporter.get('bank_name', 'HDFC Bank Ltd.')}, {exporter.get('branch_address', 'Munger Branch')}"
        ws_req['P11'] = f"{exporter.get('name', 'COD XPERT')}, {exporter.get('address', '')}"
        
        # Export Classifications (Preserving and setting valid dropdown values)
        ws_req['E12'] = 'Service'
        ws_req['AD12'] = 'Regular Export'
        ws_req['F13'] = 'Internet'
        ws_req['P14'] = 'Others (Specify)'
        ws_req['E15'] = 'Others (advance payment, etc. including transfer/remittance to bank a/c maintained overseas)'
        
        country_name = client.get('country') or 'Overseas'
        ws_req['AD15'] = f"Country of Final Destination: {country_name}"
        
        invoice_date = data.get('invoice_date_dmy') or data.get('invoice_date', '')
        ws_req['D16'] = invoice_date
        ws_req['U16'] = 'Non dispatch'
        
        services_desc = data.get('services_description') or items[0].get('name', 'Software Consultancy & IT Services')
        ws_req['G19'] = services_desc
        ws_req['G20'] = amount_in_words
        ws_req['E21'] = 'NO'  # Third party: NO
        
        # Table of Invoices (Row 25 to 28)
        client_full = f"{client.get('name', '')}, {client.get('address', '')}".strip(', ')
        client_country = country_name
        invoice_no = data.get('invoice_number', '')
        lut_num = exporter.get('lut_number', '')
        remarks = f"Export of Services under LUT No. {lut_num} without payment of IGST" if lut_num else "Export of Services under LUT without payment of IGST"
        
        start_row = 25
        max_rows = 4  # rows 25, 26, 27, 28
        for idx, item in enumerate(items[:max_rows]):
            curr_row = start_row + idx
            ws_req[f'A{curr_row}'] = idx + 1
            ws_req[f'B{curr_row}'] = client_full
            ws_req[f'E{curr_row}'] = client_country
            ws_req[f'J{curr_row}'] = invoice_no
            ws_req[f'O{curr_row}'] = invoice_date
            ws_req[f'T{curr_row}'] = currency
            item_amt = float(item.get('amount', grand_total))
            ws_req[f'U{curr_row}'] = item_amt
            ws_req[f'X{curr_row}'] = item_amt
            ws_req[f'AB{curr_row}'] = 'N/A'
            ws_req[f'AE{curr_row}'] = item.get('name') or services_desc
            ws_req[f'AH{curr_row}'] = item.get('sac_code') or data.get('sac_code', '998314')
            ws_req[f'AK{curr_row}'] = remarks

        # Footer Signoff
        ws_req['Q29'] = account_no
        ws_req['M30'] = exporter.get('phone', '')
        ws_req['AE30'] = exporter.get('email', '')
        ws_req['B35'] = f"For {exporter.get('name', 'COD XPERT')}"

        # Re-attach and ensure standard Excel dropdown DataValidations on REQUEST LETTER
        from openpyxl.worksheet.datavalidation import DataValidation
        req_dropdowns = [
            ('list', 'Sheet2!$B$2:$B$3', 'E12:M12'),      # Service / Software
            ('list', 'Sheet2!$G$1:$G$2', 'AD12:AK12'),   # Project Export / Regular Export
            ('list', 'Sheet2!$E$1:$E$6', 'F13:L13'),     # Mode of Transport (Internet, Air, etc.)
            ('list', 'Sheet2!$A$1:$A$5', 'P14'),         # Category of Exporter (Others, etc.)
            ('list', 'Sheet2!$C$1:$C$3', 'E15:N15'),     # Mode of Realisation
            ('list', 'Sheet2!$I$1:$I$3', 'U16:AA18'),    # Dispatch Indicator (Non dispatch)
            ('list', 'Sheet2!$F$1:$F$2', 'E21:G21'),     # Third Party (YES / NO)
        ]
        for dv_type, formula, cell_range in req_dropdowns:
            dv = DataValidation(type=dv_type, formula1=formula, allow_blank=True)
            ws_req.add_data_validation(dv)
            dv.add(cell_range)

    # =========================================================================
    # 2. SHEET: EDF Annexure
    # =========================================================================
    if 'EDF Annexure' in wb.sheetnames:
        ws_edf = wb['EDF Annexure']
        ws_edf['A5'] = exporter.get('account_number', '')
        ws_edf['C5'] = exporter.get('iec', 'As per Bank records')
        ws_edf['D5'] = exporter.get('gstin', 'As per Bank records')
        ws_edf['E5'] = exporter.get('pan', 'As per Bank records')
        ws_edf['G5'] = exporter.get('address', 'As per Bank records')
        ws_edf['H5'] = exporter.get('name', 'As per Bank records')
        
        # Row 11
        ws_edf['A11'] = '-'
        ws_edf['B11'] = 'service'
        ws_edf['C11'] = 'Internet'
        ws_edf['D11'] = 'Others (Specify)'
        ws_edf['E11'] = 'Others (advance payment, etc. including transfer/remittance to bank a/c maintained overseas)'
        ws_edf['F11'] = 'Regular Export'
        ws_edf['H11'] = 'Non dispatch'
        ws_edf['I11'] = services_desc
        ws_edf['J11'] = amount_in_words
        ws_edf['K11'] = client.get('name', '')
        ws_edf['L11'] = client.get('address', '')
        ws_edf['M11'] = country_name

        # Re-attach and ensure dropdown DataValidations on EDF Annexure
        from openpyxl.worksheet.datavalidation import DataValidation
        edf_dropdowns = [
            ('list', "'Format Sheet'!$B$13:$B$14", 'B11:B33'), # service / software
            ('list', "'Format Sheet'!$A$2:$A$7", 'C11:C33'),   # Internet, Air, etc.
            ('list', "'Format Sheet'!$C$2:$C$6", 'D11:D33'),   # Category
            ('list', "'Format Sheet'!$E$2:$E$4", 'E11:E33'),   # Mode of Realisation
            ('list', "'Format Sheet'!$G$12:$G$13", 'F11:F33'), # Regular Export / Project Export
            ('list', "'Format Sheet'!$D$16:$D$18", 'H11:H33'), # Non dispatch / Bank / Exporter
        ]
        for dv_type, formula, cell_range in edf_dropdowns:
            dv = DataValidation(type=dv_type, formula1=formula, allow_blank=True)
            ws_edf.add_data_validation(dv)
            dv.add(cell_range)

    # =========================================================================
    # 3. SHEET: Invoice Details 1
    # =========================================================================
    inv_sheet_name = 'Invoice Details 1 ' if 'Invoice Details 1 ' in wb.sheetnames else ('Invoice Details 1' if 'Invoice Details 1' in wb.sheetnames else None)
    if inv_sheet_name:
        ws_inv = wb[inv_sheet_name]
        start_r = 3
        for idx, item in enumerate(items[:5]):
            r = start_r + idx
            item_amt = float(item.get('amount', grand_total))
            ws_inv[f'A{r}'] = invoice_no
            ws_inv[f'B{r}'] = invoice_date
            ws_inv[f'C{r}'] = item_amt
            ws_inv[f'D{r}'] = currency
            ws_inv[f'E{r}'] = item_amt
            ws_inv[f'F{r}'] = 'N/A'
            ws_inv[f'G{r}'] = 'N/A'
            ws_inv[f'H{r}'] = item.get('name') or services_desc
            ws_inv[f'I{r}'] = item.get('sac_code') or data.get('sac_code', '998314')
            ws_inv[f'J{r}'] = remarks

    # Ensure parent directory exists and save
    os.makedirs(os.path.dirname(os.path.abspath(output_xlsx_path)), exist_ok=True)
    wb.save(output_xlsx_path)
    print(f"SUCCESS: {output_xlsx_path}")

if __name__ == '__main__':
    main()
